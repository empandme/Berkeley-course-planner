"""Turn a validated plan.json dict into the student's openpyxl Workbook.

build_workbook(plan) assumes ``plan`` already passed
plan_validation.validate_plan; it does no checking of its own.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# SECTION_ORDER and COLUMNS live in cell_text (no openpyxl) and are re-exported
# here, so ``from workbook import COLUMNS, SECTION_ORDER`` keeps working.
from cell_text import COLUMNS, SECTION_ORDER, row_cells

_COLUMN_WIDTHS = [7, 28, 14, 30, 7, 20, 18, 22, 24, 50, 20, 22, 28, 28]

REQUIREMENTS_COLUMNS = ["Group", "Requirement", "Status", "Satisfied by", "Source link"]
_REQUIREMENTS_WIDTHS = [12, 44, 16, 36, 60]

_EXCLUDED_COLUMNS = ["Course", "Slot", "Reason"]
_RUN_INFO_WIDTHS = [40, 32, 50]

_HEADER_ROW = 4
_FIRST_SECTION_ROW = _HEADER_ROW + 1
_EMPTY_SECTION_TEXT = "No candidates in this section."
_EXCLUDED_POINTER = " See Run Info → Excluded."  # only when Excluded has entries

_GREEN = PatternFill("solid", start_color="C6EFCE", end_color="C6EFCE")
_AMBER = PatternFill("solid", start_color="FFEB9C", end_color="FFEB9C")
_GREY = PatternFill("solid", start_color="D9D9D9", end_color="D9D9D9")

# "Interest match" is deliberately absent: it gets no fill.
_LABEL_FILLS = {
    "Recommended (reviews)": _GREEN,
    "Recommended (grades)": _GREEN,
    "Below bar (reviews)": _AMBER,
    "Below bar (grades)": _AMBER,
    "No data": _GREY,
}

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", start_color="305496", end_color="305496")
_SECTION_FONT = Font(bold=True)
_SECTION_FILL = PatternFill("solid", start_color="D9E1F2", end_color="D9E1F2")
_LINK_FONT = Font(color="0563C1", underline="single")
_NOTE_FONT = Font(italic=True, color="595959")

_WRAP = Alignment(wrap_text=True, vertical="top")
_HEADER_ALIGNMENT = Alignment(wrap_text=True, vertical="center")


def _put(ws, r: int, col: int, value):
    """Write value to a cell. A string always stays a string: openpyxl would
    otherwise turn "=..." into a formula and "#N/A" into an error value, and
    review notes and titles come from web pages."""
    cell = ws.cell(r, col, value)
    if isinstance(value, str):
        cell.data_type = "s"
    return cell


def _write_header_row(ws, r: int, headers: list) -> None:
    for col, header in enumerate(headers, start=1):
        cell = _put(ws, r, col, header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGNMENT


def _set_widths(ws, widths: list) -> None:
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width


def _column_letter(header: str) -> str:
    return get_column_letter(COLUMNS.index(header) + 1)


def _write_candidate(ws, r: int, row: dict) -> None:
    for col, (text, url) in enumerate(row_cells(row), start=1):
        cell = _put(ws, r, col, text if text != "" else None)
        cell.alignment = _WRAP
        if url:
            cell.hyperlink = url
            cell.font = _LINK_FONT
    fill = _LABEL_FILLS.get(row["label"])
    if fill is not None:
        ws.cell(r, COLUMNS.index("Label") + 1).fill = fill


def _write_section_title(ws, r: int, title: str) -> None:
    # Styled across A:N, not merged: Excel refuses to sort a filtered range
    # that contains merged cells.
    for col in range(1, len(COLUMNS) + 1):
        cell = ws.cell(r, col)
        cell.font = _SECTION_FONT
        cell.fill = _SECTION_FILL
    _put(ws, r, 1, title)


def _build_candidates(ws, plan: dict) -> None:
    ws.title = "Candidates"

    pick, units = _column_letter("Pick?"), _column_letter("Units")
    _put(ws, 1, 1, "Unit target")
    _put(ws, 1, 2, plan["run"]["unit_target"])
    _put(ws, 2, 1, "Units picked")
    ws["B2"] = f'=SUMIF({pick}:{pick},"Y",{units}:{units})'  # the one formula
    ws["A1"].font = ws["A2"].font = _SECTION_FONT

    _write_header_row(ws, _HEADER_ROW, COLUMNS)
    _set_widths(ws, _COLUMN_WIDTHS)

    r = _FIRST_SECTION_ROW
    for key, title in SECTION_ORDER:
        _write_section_title(ws, r, title)
        r += 1
        rows = plan["sections"][key]
        if not rows:
            text = _EMPTY_SECTION_TEXT + (_EXCLUDED_POINTER if plan["excluded"] else "")
            note = _put(ws, r, 2, text)
            note.font = _NOTE_FONT
            r += 1
        for row in rows:
            _write_candidate(ws, r, row)
            r += 1

    ws.freeze_panes = f"A{_HEADER_ROW + 1}"
    ws.auto_filter.ref = f"A{_HEADER_ROW}:{get_column_letter(len(COLUMNS))}{ws.max_row}"


def _build_requirements(ws, plan: dict) -> None:
    ws.title = "Requirements"
    _write_header_row(ws, 1, REQUIREMENTS_COLUMNS)
    _set_widths(ws, _REQUIREMENTS_WIDTHS)

    for r, item in enumerate(plan["requirements"], start=2):
        values = [
            item["group"],
            item["requirement"],
            item["status"],
            item["satisfied_by"] or None,
            item["source_url"],
        ]
        for col, value in enumerate(values, start=1):
            cell = _put(ws, r, col, value)
            cell.alignment = _WRAP
        link = ws.cell(r, len(REQUIREMENTS_COLUMNS))
        link.hyperlink = item["source_url"]
        link.font = _LINK_FONT

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(REQUIREMENTS_COLUMNS))}{ws.max_row}"


def _write_heading(ws, r: int, text: str) -> None:
    _put(ws, r, 1, text).font = _SECTION_FONT


def _build_run_info(ws, plan: dict) -> None:
    ws.title = "Run Info"
    _set_widths(ws, _RUN_INFO_WIDTHS)
    run = plan["run"]

    metadata = [
        ("Run date", run["date"]),
        ("Semester", run["semester"]),
        ("Major", run["major"]),
        ("College", run["college"]),
        ("Unit target", run["unit_target"]),
        ("Unit default applied", "Yes" if run["unit_default_applied"] else "No"),
        ("Delivery", run["delivery"]),
    ]
    r = 1
    for key, value in metadata:
        _put(ws, r, 1, key)
        _put(ws, r, 2, value)
        r += 1

    r += 1  # blank row between blocks
    _write_heading(ws, r, "Catalog years")
    r += 1
    if run["catalog_years"]:
        for catalog_file, year in run["catalog_years"].items():
            _put(ws, r, 1, catalog_file)
            _put(ws, r, 2, year)
            r += 1
    else:
        _put(ws, r, 1, "None")
        r += 1

    r += 1  # blank row between blocks
    _write_heading(ws, r, "Excluded")
    r += 1
    _write_header_row(ws, r, _EXCLUDED_COLUMNS)
    r += 1
    if plan["excluded"]:
        for item in plan["excluded"]:
            for col, value in enumerate((item["course_number"], item["slot"], item["reason"]), start=1):
                _put(ws, r, col, value).alignment = _WRAP
            r += 1
    else:
        _put(ws, r, 1, "None")


def build_workbook(plan: dict) -> Workbook:
    """Build the workbook for a plan that already passed validate_plan."""
    wb = Workbook()
    _build_candidates(wb.active, plan)
    _build_requirements(wb.create_sheet(), plan)
    _build_run_info(wb.create_sheet(), plan)
    return wb
